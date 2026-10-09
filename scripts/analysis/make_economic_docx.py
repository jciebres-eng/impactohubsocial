#!/usr/bin/env python3
"""Converte RELATORIO_EXECUTIVO_ECONOMIA_v0300.md em DOCX (python-docx) e insere os gráficos g1–g9 após a seção 5."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
D = ROOT / "docs" / "analysis" / "economia_v0300"
MD = (D / "RELATORIO_EXECUTIVO_ECONOMIA_v0300.md").read_text(encoding="utf-8").splitlines()
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else D / "IMPACTO_v0300_RELATORIO_EXECUTIVO_ECONOMIA.docx"

CHARTS = [
    ("g1_receita_por_fonte_atual.png", "Gráfico 1a — Receita mensal por fonte (desenho ATUAL, cenário base)"),
    ("g1_receita_por_fonte_hibrido.png", "Gráfico 1b — Receita mensal por fonte (desenho HÍBRIDO, cenário base)"),
    ("g2_receita_acumulada.png", "Gráfico 2 — Receita acumulada"),
    ("g3_mrr_arr.png", "Gráfico 3 — MRR e ARR de software (só no HÍBRIDO)"),
    ("g4_gmv_vs_receita.png", "Gráfico 4 — GMV versus receita da plataforma (GMV não é receita)"),
    ("g5_ebitda.png", "Gráfico 5 — EBITDA e margem"),
    ("g6_clientes.png", "Gráfico 6 — Clientes institucionais e financiadores"),
    ("g7_caixa.png", "Gráfico 7 — Caixa e necessidade de capital"),
    ("g8_cenarios_marcos.png", "Gráfico 8 — Comparação dos três cenários nos marcos"),
    ("g9_sensibilidade.png", "Gráfico 9 — Sensibilidade (Base híbrido)"),
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.6))
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(9.5)

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def runs(p, text: str, size: float | None = None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)


def table(rows: list[list[str]]):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            p = cells[j].paragraphs[0]
            runs(p, txt, 7.5)
            if i == 0:
                for r in p.runs:
                    r.bold = True
    doc.add_paragraph()


i = 0
quote_buf: list[str] = []
while i < len(MD):
    line = MD[i]
    if line.startswith("|") and i + 1 < len(MD) and re.match(r"^\|[\s:|-]+\|$", MD[i + 1]):
        rows = []
        while i < len(MD) and MD[i].startswith("|"):
            if not re.match(r"^\|[\s:|-]+\|$", MD[i]):
                rows.append([c.strip() for c in MD[i].strip().strip("|").split("|")])
            i += 1
        table(rows)
        continue
    if line.startswith("> "):
        quote_buf.append(line[2:])
        i += 1
        if i >= len(MD) or not MD[i].startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            runs(p, " ".join(quote_buf))
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x5A, 0x32, 0x00)
            quote_buf = []
        continue
    m = re.match(r"^(#{1,3}) (.*)", line)
    if m:
        level = len(m.group(1))
        h = doc.add_heading(m.group(2), level=0 if level == 1 else level)
        if level == 2 and m.group(2).startswith("6."):
            # gráficos entram antes da seção 6 (após a projeção)
            pass
        i += 1
        continue
    if re.match(r"^\d+\. ", line):
        p = doc.add_paragraph(style="List Number")
        runs(p, re.sub(r"^\d+\. ", "", line))
        i += 1
        continue
    if line.strip() == "":
        i += 1
        continue
    # parágrafo: junta linhas contínuas
    buf = [line]
    i += 1
    while i < len(MD) and MD[i].strip() and not MD[i].startswith(("|", "#", "> ")) and not re.match(r"^\d+\. ", MD[i]):
        buf.append(MD[i])
        i += 1
    p = doc.add_paragraph()
    runs(p, " ".join(buf))
    # inserir gráficos logo após a "Leitura honesta" da seção 5
    if buf[0].startswith("Modelagem separada que o pedido exige"):
        doc.add_heading("Gráficos (FASE 5) — valores em R$; cenário base salvo indicação", level=2)
        for fn, cap in CHARTS:
            doc.add_picture(str(D / fn), width=Cm(24))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(c, cap, 8.5)
            for r in c.runs:
                r.italic = True

doc.save(OUT)
print("ok", OUT, OUT.stat().st_size, "bytes")
